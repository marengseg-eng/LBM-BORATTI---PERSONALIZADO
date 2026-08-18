"""
lib_docx.py — utilitários para editar PPP e LTCAT (.docx) preservando a estrutura.

Encapsula o que foi validado na prática:
  - inserir uma linha de risco numa tabela clonando o layout/mesclagens de uma linha
    existente (evita quebrar mesclagens de célula, comum no formulário do PPP);
  - travar linhas de tabela para não se partirem entre páginas (cantSplit);
  - manter o bloco de assinatura unido (keep_together / keep_with_next).

Uso típico dentro da skill:

    from docx import Document
    import lib_docx as L

    doc = Document("PPP.docx")
    tab = doc.tables[0]                     # tabela do campo 15 do PPP
    idx_ruido = L.indice_linha_por_texto(tab, "Ruído")   # acha a linha modelo
    L.clonar_linha_risco(
        tab, idx_ruido,
        {1: "Biológico",
         2: "Agentes biológicos (bactérias, vírus, fungos e parasitas) - contato "
            "permanente com bovinos, secreções e dejetos em currais/estábulos",
         3: "Qualitativo",
         4: "Qualitativa - NR-15, Anexo 14",
         5: "NA", 6: "N", 7: "NI"},
        cant_split=True)
    doc.save("PPP.docx")

IMPORTANTE: os índices de coluna são a ORDEM das células <w:tc> da linha modelo,
não os números de campo do formulário. Rode `mapear_linha(tab, idx)` para vê-los
antes de montar o dicionário de valores.
"""

import copy
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def mapear_linha(tabela, indice_linha):
    """Retorna [(i, texto)] de cada célula <w:tc> da linha — use para descobrir
    quais índices preencher antes de clonar."""
    tr = tabela.rows[indice_linha]._tr
    out = []
    for i, tc in enumerate(tr.findall(qn('w:tc'))):
        txt = ' '.join(n.text or '' for n in tc.iter(qn('w:t')) if n.text).strip()
        out.append((i, txt))
    return out


def indice_linha_por_texto(tabela, texto):
    """Índice da primeira linha cujo conteúdo contém `texto` (case-insensitive).
    Útil para localizar a linha-modelo (ex.: a linha do 'Ruído') sem fixar número."""
    alvo = texto.lower()
    for ri, row in enumerate(tabela.rows):
        conteudo = ' '.join(c.text for c in row.cells).lower()
        if alvo in conteudo:
            return ri
    raise ValueError(f"Linha com texto {texto!r} não encontrada.")


def _set_texto_tc(tc, texto):
    """Escreve `texto` na primeira run da célula e zera as demais, preservando a
    formatação da primeira run (fonte, tamanho, negrito)."""
    wts = [n for n in tc.iter(qn('w:t'))]
    if wts:
        wts[0].text = texto
        for extra in wts[1:]:
            extra.text = ''


def set_cant_split(tr):
    """Marca a linha (<w:tr>) como 'não dividir entre páginas'."""
    trPr = tr.find(qn('w:trPr'))
    if trPr is None:
        trPr = OxmlElement('w:trPr')
        tr.insert(0, trPr)
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))


def clonar_linha_risco(tabela, indice_modelo, valores, cant_split=True):
    """Clona a linha `indice_modelo` (mantendo mesclagens/estilo) e a insere logo
    abaixo, preenchendo as células segundo `valores` = {indice_tc: texto}.
    Índices ausentes em `valores` mantêm o conteúdo do modelo (ex.: a coluna de
    período, que costuma ser igual)."""
    tr_modelo = tabela.rows[indice_modelo]._tr
    nova = copy.deepcopy(tr_modelo)
    tcs = nova.findall(qn('w:tc'))
    for i, tc in enumerate(tcs):
        if i in valores:
            _set_texto_tc(tc, valores[i])
    if cant_split:
        set_cant_split(nova)
        set_cant_split(tr_modelo)   # trava também a linha-modelo, por simetria
    tr_modelo.addnext(nova)
    return nova


def travar_assinatura(doc, marcadores):
    """Mantém o bloco de assinatura unido, evitando que nome/CREA se separem entre
    páginas (o defeito clássico de o CREA cair sozinho no topo da página seguinte).
    `marcadores` = lista de trechos iniciais dos parágrafos da assinatura
    (ex.: ["Marcelo Luis Boratti Melo", "CREA"]). Aplica keep_together + keep_with_next."""
    achou = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if any(t.startswith(m) for m in marcadores):
            pf = p.paragraph_format
            pf.keep_together = True
            pf.keep_with_next = True
            achou.append(t[:40])
    return achou


def remover_paragrafos_vazios_excedentes(doc, apos_marcador, manter=2):
    """Reduz sequências longas de parágrafos em branco que empurram a assinatura
    para a página seguinte. Começa a contar após o parágrafo que inicia com
    `apos_marcador` e mantém no máximo `manter` linhas em branco seguidas."""
    ps = doc.paragraphs
    inicio = None
    for i, p in enumerate(ps):
        if p.text.strip().startswith(apos_marcador):
            inicio = i
            break
    if inicio is None:
        return 0
    removidos = 0
    seguidos = 0
    for p in ps[inicio + 1:]:
        if p.text.strip() == '':
            seguidos += 1
            if seguidos > manter:
                el = p._p
                el.getparent().remove(el)
                removidos += 1
        else:
            seguidos = 0
    return removidos
